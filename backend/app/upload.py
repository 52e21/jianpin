"""简历文件提取：支持 PDF / DOCX（纯本地处理，零 LLM）。"""
import fitz  # PyMuPDF
from docx import Document
from fastapi import HTTPException

# 文件大小上限 5MB
MAX_FILE_SIZE = 5 * 1024 * 1024


def extract_pdf(content: bytes) -> str:
    doc = fitz.open(stream=content, filetype="pdf")
    text = ""
    for page in doc:
        text += page.get_text()
    return text.strip()


def extract_docx(content: bytes) -> str:
    from io import BytesIO

    doc = Document(BytesIO(content))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    # 也尝试表格文本
    table_text = []
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    table_text.append(cell.text.strip())
    return "\n".join(paragraphs + table_text).strip()


def extract_resume(filename: str, content: bytes) -> str:
    """按后缀分派提取逻辑。

    第 15 步：与前端契约对齐 —— 前端 accept 已收紧为 .pdf/.docx。
    对旧版 .doc 给出可操作的提示（python-docx 只支持 OOXML 的 .docx，
    老的 .doc 二进制格式需要额外转换工具，本项目不引入）。
    """
    name = (filename or "").lower()
    if name.endswith(".pdf"):
        return extract_pdf(content)
    elif name.endswith(".docx"):
        return extract_docx(content)
    elif name.endswith(".doc"):
        raise HTTPException(
            status_code=400,
            detail="不支持旧版 .doc 格式，请用 Word「另存为 .docx」后重新上传",
        )
    else:
        raise HTTPException(status_code=400, detail="仅支持 PDF 或 DOCX 文件")
